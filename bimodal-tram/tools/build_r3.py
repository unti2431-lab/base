# -*- coding: utf-8 -*-
"""r3_data.py → BIMODAL_R3_CUES.xlsx, R3_CUES.md, R3_MASTER.md, prompts/FLOW_PROMPTS.md, script/*.txt

    python3 build_r3.py

타임코드는 음절 수 기반의 명목값이다. 녹음 후 실측 길이로 시작·종료 초만 고치면
엑셀의 길이·프레임 열은 수식으로 다시 계산된다.
"""
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import r3_data as D

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def syllables(text):
    return len(re.findall(r"[가-힣0-9A-Za-z]", text))


def assign_times(rows, rate, pause, target, tail=1.0):
    """음절 수 비례로 길이를 잡고 목표 길이에 맞게 스케일한 뒤 0.5초로 반올림한다.
    마지막 샷에는 발화 뒤 환경음 tail초를 따로 더한다."""
    raw = [syllables(r["vo"]) / rate + pause + r.get("hold", 0) for r in rows]
    raw[-1] += tail * sum(raw) / (target - tail)
    scale = target / sum(raw)
    t = 0.0
    for i, (r, d) in enumerate(zip(rows, raw)):
        r["start"] = t
        end = target if i == len(rows) - 1 else round((t + d * scale) * 2) / 2
        r["end"] = end
        r["syl"] = syllables(r["vo"])
        r["sps"] = r["syl"] / (end - t - (tail if i == len(rows) - 1 else 0))
        t = end
    return rows


def tc(sec):
    m, s = divmod(sec, 60)
    return f"{int(m):02d}:{s:04.1f}" if s % 1 else f"{int(m):02d}:{int(s):02d}"


SHORT = assign_times([dict(r) for r in D.SHORT], D.SHORT_RATE, D.SHORT_PAUSE, D.SHORT_TARGET)
LONG = assign_times([dict(r) for r in D.LONG], D.LONG_RATE, D.LONG_PAUSE, D.LONG_TARGET)
SRC = {s[0]: s for s in D.SOURCES}
SEC = dict(D.LONG_SECTIONS)


def src_urls(ids):
    return "\n".join(SRC[i.strip()][3] for i in ids.split(";") if i.strip() in SRC)


# --------------------------------------------------------------------------- xlsx
HEAD = PatternFill("solid", fgColor="1F2A33")
KEY = PatternFill("solid", fgColor="FFF4D6")
WARN = PatternFill("solid", fgColor="FDE2DD")
THIN = Side(style="thin", color="C9CED2")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet(wb, name, title, note, header, rows, widths, key_rows=()):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = note
    ws["A3"].font = Font(italic=True, color="555555")
    for c, h in enumerate(header, 1):
        cell = ws.cell(row=5, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = WRAP
        cell.border = BOX
    for r, row in enumerate(rows, 6):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = WRAP
            cell.border = BOX
            if (r - 6) in key_rows:
                cell.fill = KEY
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B6"
    return ws


def build_xlsx(path):
    wb = Workbook()
    ov = wb.active
    ov.title = "Overview"
    lines = [
        (f"바이모달 트램 | 대본·큐시트·Flow 생성표 R3", None),
        (f"{D.DATE} · 흥행(생활 훅·반전·수미상관) × 정확성(원리/기록/원인/비용 분리) · Google Flow 단독 생성", None),
        (None, None),
        ("항목", "설정"),
        ("기준 FPS", D.FPS),
        ("쇼츠 목표(초)", D.SHORT_TARGET),
        ("롱폼 목표(초)", D.LONG_TARGET),
        ("쇼츠 규격", "9:16 · 1080×1920 · Veo 네이티브 세로 생성"),
        ("롱폼 규격", "16:9 · 1920×1080 · 쇼츠 키프레임을 참조로 16:9 쌍둥이 재생성"),
        ("쇼츠 제목", D.SHORT_TITLE),
        ("롱폼 제목", D.LONG_TITLE),
        ("타깃", "45~75세 남성 · 교통·기계 관심층 (채널 실측 데이터 미조회, 편집 판단)"),
        ("낭독 기준", f"쇼츠 {D.SHORT_RATE}음절/초 · 롱폼 {D.LONG_RATE}음절/초 + 컷 사이 쉼 (명목, 녹음 후 재조정)"),
        ("프레임 규약", "0-based / end-exclusive"),
        (None, None),
        ("검산", "샷 수 / 합계(초) / 차이(초)"),
        ("Short_R3", None),
        ("Long_R3", None),
        (None, None),
        ("읽는 순서", "Overview → Short_R3 / Long_R3 → Flow_KF → Flow_Clips → Sources / Claims → QA"),
        ("핵심 컷", f"S09–S11 ({SHORT[10]['end'] - SHORT[8]['start']:g}초): 바퀴가 떠난 뒤 카메라가 남는다 → 얕은 잔류 → 쇠자 밑 빛 틈. 이 구간이 안 읽히면 나머지를 만들지 않는다."),
        ("상태", "대본·명목 큐·생성 프롬프트 작성 완료. Flow 생성·TTS·편집·렌더는 미수행. 근거 F01·F12–F14는 녹음 전 원문 재대조 필요."),
    ]
    for r, (a, b) in enumerate(lines, 1):
        ov.cell(row=r, column=1, value=a)
        ov.cell(row=r, column=2, value=b)
    ov["A1"].font = Font(bold=True, size=14)
    ov["B17"] = f"=COUNTA(Short_R3!A6:A{5 + len(SHORT)})"
    ov["C17"] = f"=SUM(Short_R3!D6:D{5 + len(SHORT)})"
    ov["D17"] = "=C17-B6"
    ov["B18"] = f"=COUNTA(Long_R3!A6:A{5 + len(LONG)})"
    ov["C18"] = f"=SUM(Long_R3!E6:E{5 + len(LONG)})"
    ov["D18"] = "=C18-B7"
    ov.column_dimensions["A"].width = 16
    ov.column_dimensions["B"].width = 110

    # Short
    rows = []
    for i, r in enumerate(SHORT):
        n = 6 + i
        rows.append([r["id"], r["start"], r["end"], f"=C{n}-B{n}", f"=B{n}*Overview!$B$5", f"=C{n}*Overview!$B$5",
                     r["vo"], r["syl"], round(r["sps"], 2), r["role"], r["visual"], r["camera"], r["sfx"],
                     r["kf"], r["mode"], r["check"], r["src"], src_urls(r["src"]), "생성 전"])
    key = [i for i, r in enumerate(SHORT) if "★" in r["role"]]
    sheet(wb, "Short_R3", f"SHORT 58 | {D.SHORT_TITLE}",
          "시작·종료 초만 수정 · 길이·프레임은 수식 · ★ = 핵심 구간 · 음절/초 6.5 초과는 빠름",
          ["샷", "시작(초)", "종료(초)", "길이(초)", "시작F", "종료F(미포함)", "내레이션", "음절", "음절/초", "역할",
           "화면", "카메라", "SFX·음악", "키프레임", "Flow 생성", "검수 핵심", "근거", "근거 URL", "상태"],
          rows, [6, 8, 8, 8, 8, 9, 40, 6, 7, 12, 46, 22, 26, 14, 26, 32, 10, 40, 10], key)

    # Long
    rows = []
    for i, r in enumerate(LONG):
        n = 6 + i
        rows.append([r["id"], f"{r['sec']} {SEC[r['sec']]}", r["start"], r["end"], f"=D{n}-C{n}",
                     f"=C{n}*Overview!$B$5", f"=D{n}*Overview!$B$5", r["vo"], r["syl"], round(r["sps"], 2),
                     r["visual"], r["sfx"], r["kf"], r["mode"], r["check"], r["src"], src_urls(r["src"]), "생성 전"])
    key = [i for i, r in enumerate(LONG) if r["id"] in KEY_LONG]
    sheet(wb, "Long_R3", f"LONG 10 | {D.LONG_TITLE}",
          "시작·종료 초만 수정 · 한 주샷 안에서 관찰 대상 2–3개를 순차 공개 · Flow 클립은 8초 생성 후 트림",
          ["샷", "장", "시작(초)", "종료(초)", "길이(초)", "시작F", "종료F(미포함)", "내레이션", "음절", "음절/초",
           "화면", "SFX·음악", "키프레임", "Flow 생성", "검수 핵심", "근거", "근거 URL", "상태"],
          rows, [6, 22, 8, 8, 8, 8, 9, 50, 6, 7, 44, 24, 16, 28, 30, 10, 40, 10], key)

    # Flow KF
    rows = []
    for k in sorted(D.KEYFRAMES, key=lambda k: k["order"]):
        twin = "필요" if ("9:16" in k["prompt"] and "C" in k["used"]) else ""
        rows.append([k["order"], k["id"], k["used"], k["pair"] or "—", k["prompt"], k["locks"], twin, "생성 전"])
    sheet(wb, "Flow_KF", "FLOW KEYFRAMES | 생성 순서 = 위험한 것부터",
          "LOCK 본문은 prompts/LOCKS.md · 선행 KF가 있으면 그 이미지를 확정·첨부한 뒤 순차 생성(병렬 금지) · 16:9 쌍둥이: 롱폼용으로 같은 장면을 가로로 재생성",
          ["순서", "KF", "사용 샷", "선행·참조", "프롬프트(EN, 피사체 선언 먼저)", "LOCK 태그(끝에 붙임)", "16:9 쌍둥이", "상태"],
          rows, [6, 9, 18, 10, 90, 34, 10, 10])

    # Flow clips
    rows = []
    for r in SHORT:
        rows.append([f"SH-{r['id']}", r["id"], r["kf"], r["mode"], "9:16 · 8초 생성 → 트림",
                     D.SHORT_MOTION.get(r["id"], "(TRIM/편집 합성 — 신규 생성 없음)"), "Veo 오디오는 앰비언스 참고만, VO 트랙과 분리"])
    for r in LONG:
        rows.append([f"LF-{r['id']}", r["id"], r["kf"], r["mode"], "16:9", r["visual"], ""])
    sheet(wb, "Flow_Clips", "FLOW CLIPS | 샷별 생성 방식",
          "F2V = Frames to Video(시작/끝 프레임) · ING = Ingredients to Video(참조 최대 3장) · TRIM = 기존 클립 재사용 · 8초×n = 같은 시작 프레임으로 n개 생성 또는 Extend",
          ["클립", "샷", "키프레임", "방식", "규격", "움직임 지시", "비고"], rows, [10, 6, 18, 30, 16, 90, 30])

    # Sources
    sheet(wb, "Sources", "SOURCES | 원문 · 작성일 · 사용 범위 · 확인 상태", "열람 기준 " + D.DATE,
          ["ID", "자료", "작성일", "URL", "사용", "제한", "확인 상태"], [list(s) for s in D.SOURCES],
          [6, 46, 18, 50, 46, 46, 30])

    # Claims
    sheet(wb, "Claims", "CLAIMS | 원문 → R2 → R3", "R3은 R2의 사실 교정을 모두 유지하고 훅·구조·연출을 바꿨다",
          ["항목", "원문", "R2 처리", "R3 처리", "근거"], CLAIMS, [16, 36, 36, 50, 12])

    # QA
    sheet(wb, "QA", "QA | 제작 조건 · 확인하지 않은 범위", "", ["#", "조건"],
          [[i + 1, q] for i, q in enumerate(QA)], [5, 140])

    wb.save(path)


KEY_LONG = ("C16", "C21", "C22", "C26", "C27")

CLAIMS = [
    ["훅", "900억짜리 차가 폐차장으로", "삭제. 도로 흔적에서 시작", "누구나 본 '정류장 앞 콘크리트 판'에서 시작. 생활 관찰 → 원리 → 바이모달 트램", "F12; F13"],
    ["금액", "900억 / 수백억 콘크리트", "삭제", "유지(삭제). 롱폼 L10에서 '비교 조건'만 제시", "—"],
    ["폐차 서사", "만들자마자 폐차", "삭제, 2018 청라 4대 투입 기록 제시", "유지. 롱폼 C29에서 '결이 다르다'로 정리", "F03"],
    ["유도 원리", "레이저가 자석을 스캔", "자기장을 읽는 센서", "유지 + '이정표에 가깝다' 비유, 센서·마커 비접촉 화면", "F01; F02"],
    ["정밀도 수치", "0.1cm / 0.001%", "삭제", "유지(삭제). '거의 같은 줄'", "—"],
    ["사람 운전 버스", "파임 없음", "교정: 위치 분산 vs 집중", "강화: 정류장 훅 자체가 '사람이 모는 버스도 홈을 만든다'를 증명", "F08; F12"],
    ["재료", "아스팔트가 엿처럼 녹음", "골재·바인더 복합재", "유지 + 시험실 휠트래킹·절단 시편 실사 화면으로 시각화", "F06; F07"],
    ["연쇄 고장", "파임→센서 먹통→유압 파열", "삭제", "유지(삭제). 롱폼 C32에서 '연결 자료 없음' 명시", "F04"],
    ["설계자 무지", "몰랐다", "특허 반증", "유지. 쇼츠 S12–S13 반전 구간으로 승격", "F01"],
    ["경제성", "경전철이 무조건 싸다", "판정 대신 비교 조건", "유지. '세금 감시 질문'으로 타깃 관심사와 연결", "—"],
    ["신규: 정류장 콘크리트", "—", "—", "서울시 2010 공항대로 시범·2020 정류장 8곳 발표. 시 발표 성능 수치(포트홀 0건·7년/20년)는 쓰지 않음", "F12"],
    ["신규: 쇠자 측정", "—", "—", "직선자로 홈 깊이를 보는 측정 방식을 '증거 화면'으로 사용. 눈금·숫자는 화면에 넣지 않음", "F06"],
]

QA = [
    "Flow 단독 생성: 모든 실사 화면은 Flow(이미지 생성 → Frames to Video / Ingredients to Video / Extend)로 만든다. 예외는 실제 문서 서지 캡처 합성(F01·F03·F04·F05·F12)뿐이며, 이는 증거를 보여주기 위한 것이다. 생성형 가짜 문서 금지.",
    "한글: Veo는 한글을 깨뜨린다. 자막·라벨·서지 카드는 전부 편집에서 넣는다. 생성 이미지에 글자가 한 자라도 보이면 불합격.",
    "자막(45~75세 대상): 화면 높이의 6% 이상 크기, 최대 2줄, 흰 글자 + 반투명 검정 띠. 쇼츠는 하단 1/3 위쪽(UI 가림 회피). 핵심 단어 1개만 노란색.",
    "쇼츠는 9:16로 네이티브 생성한다. 롱폼은 16:9로 따로 생성하되 쇼츠 확정 키프레임을 참조로 첨부한 '16:9 쌍둥이'로 만든다. 9:16을 잘라 16:9로 쓰지 않는다(해상도·구도 손실).",
    "쌍 컷(KF09A→B, KF10A→B, KF07→B, KF13→R, KF01→KF12)은 선행 KF 확정 후 참조 첨부로 순차 생성한다. 병렬 생성 금지.",
    "잔류변형의 크기: 시험실 매크로에서도 '얕게'. 쇠자 밑 빛 틈이 보이는 정도가 상한이다. 깊은 도랑·균열·포트홀은 정확성 위반.",
    "시험 장면은 '원리 재현'이다. 실제 시험 결과·횟수·깊이를 화면에 표시하지 않는다. 쇼츠 S08–S11, 롱폼 C15–C22에 '재현' 표기를 편집에서 작게 넣는다.",
    "차량은 일반화 외형이다. 실제 차량 도색·로고·형식 재현 시도 금지. 굴절부 1개·3축 유지가 형상 검수 기준.",
    "근거 F01(특허 문장), F12(서울시 발표 게시일·문장), F13(논문 서지), F14는 이번 세션에서 원문 사이트 접근이 차단돼 검색 색인 인용으로만 확인했다. 녹음 전 원문을 직접 열어 대조한다.",
    "청라 노선에서 자동 유도 모드를 실제로 썼는지는 확인하지 않았다. 대본은 '설계됐다'로만 말한다.",
    "타깃 판단은 편집 판단이다. 채널 분석 데이터, 유지율, 클릭률을 조회하지 않았다.",
    "Veo 생성 오디오는 앰비언스 참고용. 최종 사운드는 VO + 효과음 + 환경음을 편집에서 믹스한다. 정류장·시험실의 '쉼' 구간(S09–S10, C16)은 음악을 뺀다.",
    "길이: 타임코드는 음절 수 기반의 명목값이다. 실제 녹음 길이를 재고 시작·종료 초만 고친다.",
    "설명란 고지: '도로 단면·시험실·반복 주행 장면은 원리 설명을 위한 AI 생성 재현 영상입니다. 시간과 변형은 압축·확대했으며 특정 노선의 실제 파손이나 시험 결과가 아닙니다. 역사적 사실은 기재한 문서의 해당 날짜 기준입니다.'",
    "플랫폼 AI 생성 콘텐츠 표시: 업로드 시 사실적 합성 콘텐츠 표시 항목을 켠다.",
]


# --------------------------------------------------------------------------- md / txt
def md_table(header, rows):
    out = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    for r in rows:
        out.append("| " + " | ".join(str(c).replace("|", "/").replace("\n", " ") for c in r) + " |")
    return "\n".join(out)


def build_md(path):
    s = [f"# 바이모달 트램 R3 — 큐시트 (자동 생성)", "",
         f"`tools/build_r3.py`가 `tools/r3_data.py`에서 만든다. 손으로 고치지 말고 데이터 파일을 고친 뒤 다시 빌드한다.", "",
         f"## 쇼츠 — {D.SHORT_TITLE} ({D.SHORT_TARGET}초 · 9:16 · {len(SHORT)}컷)", ""]
    s.append(md_table(["샷", "구간", "내레이션", "화면", "카메라 · SFX", "KF · Flow", "근거"],
                      [[r["id"] + (" ★" if "★" in r["role"] else ""), f"{tc(r['start'])}–{tc(r['end'])}", r["vo"], r["visual"],
                        f"{r['camera']} / {r['sfx']}", f"{r['kf']} · {r['mode']}", r["src"]] for r in SHORT]))
    s += ["", f"## 롱폼 — {D.LONG_TITLE} ({D.LONG_TARGET // 60}분 · 16:9 · {len(LONG)}주샷)", ""]
    for sid, name in D.LONG_SECTIONS:
        rs = [r for r in LONG if r["sec"] == sid]
        s += [f"### {sid} {name} — {tc(rs[0]['start'])}–{tc(rs[-1]['end'])}", ""]
        s.append(md_table(["샷", "구간", "내레이션", "화면", "KF · Flow", "검수", "근거"],
                          [[r["id"] + (" ★" if r["id"] in KEY_LONG else ""), f"{tc(r['start'])}–{tc(r['end'])}", r["vo"], r["visual"], f"{r['kf']} · {r['mode']}",
                            r["check"], r["src"]] for r in rs]))
        s.append("")
    s += ["## 키프레임 생성 순서", "",
          md_table(["순서", "KF", "사용", "선행·참조", "LOCK"],
                   [[k["order"], k["id"], k["used"], k["pair"] or "—", k["locks"]]
                    for k in sorted(D.KEYFRAMES, key=lambda k: k["order"])]), ""]
    open(path, "w").write("\n".join(s))


def build_txt():
    open(os.path.join(ROOT, "script", "SHORT_R3_VO.txt"), "w").write(
        f"[쇼츠] {D.SHORT_TITLE}\n\n" + "\n\n".join(f"({r['id']}) {r['vo']}" for r in SHORT) + "\n")
    out = [f"[롱폼] {D.LONG_TITLE}\n"]
    for sid, name in D.LONG_SECTIONS:
        out.append(f"\n## {sid} {name}\n")
        out += [f"({r['id']}) {r['vo']}\n" for r in LONG if r["sec"] == sid]
    open(os.path.join(ROOT, "script", "LONG_R3_VO.txt"), "w").write("\n".join(out))


def build_prompts(path):
    s = ["# Flow 프롬프트 — 바이모달 트램 R3 (자동 생성)", "",
         "`tools/r3_data.py`에서 만든다. LOCK 태그 자리에는 `prompts/LOCKS.md`의 본문을 그대로 붙인다.",
         "붙이는 순서: **프롬프트 본문 → 대상·장소 LOCK → STYLE LOCK(맨 끝)**. 선행 KF가 있으면 그 확정본을 참조 이미지로 첨부한다.", "",
         "## 1. 키프레임 (Flow 이미지 생성) — 생성 순서대로", ""]
    for k in sorted(D.KEYFRAMES, key=lambda k: k["order"]):
        ref = f" · **참조 첨부: {k['pair']} 확정본**" if k["pair"] else ""
        s += [f"### {k['order']:02d}. {k['id']} — {k['used']}{ref}", "", "```", k["prompt"], "", k["locks"], "```", ""]
    s += ["## 2. 쇼츠 영상 (Frames to Video) — 움직임 지시", "",
          "시작 프레임(필요하면 끝 프레임)에 해당 KF를 넣고, 아래 움직임 지시 뒤에 같은 KF의 LOCK 태그를 다시 붙인다. 8초로 생성해 편집에서 트림한다.", ""]
    for r in SHORT:
        m = D.SHORT_MOTION.get(r["id"])
        s += [f"### SH-{r['id']} — {r['kf']} · {r['mode']}", ""]
        s += ["```", m, "```", ""] if m else ["신규 생성 없음 — " + r["mode"], ""]
    s += ["## 3. 16:9 쌍둥이 (롱폼용)", "",
          "쇼츠에서 확정한 9:16 KF를 참조로 첨부하고 아래 문장 + 원래 LOCK 태그로 다시 생성한다.", "", "```",
          "Reproduce the attached reference scene exactly: same place, same objects, same light, same grade, same camera height and direction. Widen the framing sideways to a horizontal 16:9 composition; add only more of the same surroundings at the sides. Change nothing else.",
          "```", ""]
    open(path, "w").write("\n".join(s))


def build_master(path):
    t = open(os.path.join(ROOT, "tools", "master_template.md")).read()
    sh = []
    for r in SHORT:
        mark = " ★" if "★" in r["role"] else ""
        sh.append(f"> `{r['id']}{mark} {tc(r['start'])}` {r['vo']}")
    groups = [("생활 관찰", "S01", "S04"), ("같은 원리의 다른 차", "S05", "S07"), ("★ 시험실 증명", "S08", "S11"),
              ("반전과 판단 기준", "S12", "S14"), ("결론과 수미상관", "S15", "S15")]
    by = {r["id"]: r for r in SHORT}
    structure = " → ".join(f"{n}({by[a]['start']:g}–{by[b]['end']:g}초)" for n, a, b in groups)
    lg = []
    for sid, name in D.LONG_SECTIONS:
        rs = [r for r in LONG if r["sec"] == sid]
        lg += [f"#### {sid} · {tc(rs[0]['start'])}–{tc(rs[-1]['end'])} | {name}", ""]
        lg += [f"{r['vo']} <sub>{r['id']}</sub>" + "\n" for r in rs]
        lg.append(f"<sub>근거: {', '.join(sorted({x.strip() for r in rs for x in r['src'].split(';')} & set(SRC)))}</sub>\n")
    gens = 0
    for r in LONG:
        for x in re.findall(r"8초(?:×(\d))?", r["mode"]):
            gens += int(x) if x else 1
    lens = [r["end"] - r["start"] for r in LONG]
    vals = {
        "DATE": D.DATE, "SHORT_TITLE": D.SHORT_TITLE, "LONG_TITLE": D.LONG_TITLE, "SHORT_TARGET": str(D.SHORT_TARGET),
        "SHORT_N": str(len(SHORT)), "LONG_N": str(len(LONG)), "SHORT_SCRIPT": "\n>\n".join(sh),
        "LONG_SCRIPT": "\n".join(lg), "SHORT_STRUCTURE": structure,
        "QUIET_SEC": f"{by['S10']['end'] - by['S09']['start']:g}",
        "SHORT_MIN": f"{min(r['end'] - r['start'] for r in SHORT):g}", "LONG_MIN": f"{min(lens):g}", "LONG_MAX": f"{max(lens):g}",
        "LONG_GENS": str(gens), "LONG_GEN_SEC": str(gens * 8),
        "SOURCES": "\n".join(f"| {x[0]} | [{x[1]}]({x[3]}) | {x[2]} | {x[6]} |" for x in D.SOURCES),
    }
    for k, v in vals.items():
        t = t.replace("{{" + k + "}}", v)
    assert "{{" not in t, re.findall(r"{{\w+}}", t)
    open(path, "w").write(t)


def report():
    for name, rows in (("SHORT", SHORT), ("LONG", LONG)):
        tot = sum(r["syl"] for r in rows)
        fast = [f"{r['id']}({r['sps']:.1f})" for r in rows if r["sps"] > 6.5]
        print(f"{name}: {len(rows)} shots, {tot} syl, end {rows[-1]['end']}s, avg {tot / rows[-1]['end']:.2f} syl/s, "
              f"min {min(r['end'] - r['start'] for r in rows)}s, fast {fast}")
    print(sum(1 for _ in D.KEYFRAMES), "keyframes")


if __name__ == "__main__":
    build_xlsx(os.path.join(ROOT, "BIMODAL_R3_CUES.xlsx"))
    build_md(os.path.join(ROOT, "R3_CUES.md"))
    build_txt()
    build_prompts(os.path.join(ROOT, "prompts", "FLOW_PROMPTS.md"))
    build_master(os.path.join(ROOT, "R3_MASTER.md"))
    report()
